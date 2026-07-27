#!/usr/bin/env python3
"""Read first worksheet from .xlsx and print rows as JSON (for PHP when ext-zip is missing)."""
import json
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    print(json.dumps({"error": "openpyxl not installed. Run: python -m pip install openpyxl"}))
    sys.exit(1)


def main() -> None:
    if len(sys.argv) < 2:
        print(json.dumps({"error": "Usage: xlsx_sheet_rows.py <file.xlsx> [sheet_index]"}))
        sys.exit(1)

    path = sys.argv[1]
    sheet_index = int(sys.argv[2]) if len(sys.argv) > 2 else 0

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = wb.worksheets
    if sheet_index < 0 or sheet_index >= len(sheets):
        print(json.dumps({"error": f"Sheet index {sheet_index} out of range"}))
        sys.exit(1)

    ws = sheets[sheet_index]
    rows = []
    for row in ws.iter_rows(values_only=True):
        cells = []
        for cell in row:
            if cell is None:
                cells.append(None)
            else:
                cells.append(str(cell).strip() if str(cell).strip() != "" else None)
        while cells and cells[-1] is None:
            cells.pop()
        if any(c is not None and c != "" for c in cells):
            rows.append(cells)

    wb.close()
    print(json.dumps({"rows": rows}, ensure_ascii=False))


if __name__ == "__main__":
    main()
